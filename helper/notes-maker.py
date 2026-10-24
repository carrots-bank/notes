#!/usr/bin/env python3
"""Notes maker: turns a recording and documents into notes in the course format, on the person's own laptop.

Run by the 🫏 helper inside its own Python environment (venv, set up by `install.sh --maker`):
    notes-maker.py JOB_DIR
JOB_DIR holds job.json (person's secret, worker, branch/topic/title/date, engine) and files/ (the uploads).
Steps: read slides/documents -> transcribe the recording locally (Whisper, MLX) -> write notes with the person's
Claude Code (or Codex) following notes_guide.md -> encrypt and save the notes and their images to the person's
private space on the notes worker (same keys and record format as the website's space.js / annotate.js).
Progress is printed as JSON lines for the helper to show.
"""
import base64
import hashlib
import hmac
import json
import os
import re
import secrets
import shutil
import subprocess
import sys
import time
import urllib.error
import urllib.request
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
GUIDE = HERE / "notes_guide.md"
AUDIO_EXT = {".mp3", ".m4a", ".wav", ".aac", ".flac", ".ogg", ".opus", ".mp4", ".mov", ".mkv", ".webm", ".m4v", ".avi"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp"}
SALT = b"lecture-notes-annotations"


def say(**kw):
    print(json.dumps(kw), flush=True)


def step(msg):
    say(step=msg)


# ---------- keys and storage (mirrors annotate.js) ----------
class Space:
    def __init__(self, secret_hex, worker):
        from cryptography.hazmat.primitives import hashes
        from cryptography.hazmat.primitives.ciphers.aead import AESGCM
        from cryptography.hazmat.primitives.kdf.hkdf import HKDF
        secret = bytes.fromhex(secret_hex)
        d = lambda info: HKDF(hashes.SHA256(), 32, SALT, info).derive(secret)
        self.token = d(b"token").hex()
        self.aes = AESGCM(d(b"enc"))
        self.doc_key = d(b"doc")
        self.worker = worker.rstrip("/")

    def doc_id(self, path):
        return hmac.new(self.doc_key, path.encode(), hashlib.sha256).hexdigest()[:32]

    def seal(self, value):
        iv = os.urandom(12)
        ct = self.aes.encrypt(iv, json.dumps(value).encode(), None)
        return base64.b64encode(iv).decode() + "." + base64.b64encode(ct).decode()

    def open(self, data):
        iv, ct = (base64.b64decode(x) for x in data.split("."))
        return json.loads(self.aes.decrypt(iv, ct, None))

    def call(self, method, path, body=None):
        req = urllib.request.Request(f"{self.worker}/v1/doc/{self.doc_id(path)}", method=method,
                                     data=json.dumps(body).encode() if body is not None else None,
                                     headers={"authorization": "Bearer " + self.token, "content-type": "application/json",
                                              "user-agent": "NotesMaker/1"})
        try:
            with urllib.request.urlopen(req, timeout=60) as r:
                return r.status, json.loads(r.read() or b"{}")
        except urllib.error.HTTPError as e:
            body = json.loads(e.read() or b"{}")
            if e.code == 409:
                return 409, body
            raise RuntimeError(f"storage error {e.code}: {body.get('error', '')}")

    def get(self, path):
        _, j = self.call("GET", path)
        return (self.open(j["data"]) if j.get("data") else None), j.get("updated", 0)

    def put(self, path, value, merge=None):
        cur, base = self.get(path)
        if merge and cur:  # merge with what is stored now, so other devices' changes are kept
            value = merge(cur, value)
        for _ in range(5):
            st, j = self.call("PUT", path, {"data": self.seal(value), "base": base})
            if st != 409:
                return value
            if not merge:
                raise RuntimeError("conflict")
            base = j["updated"]
            value = merge(self.open(j["data"]), value)
        raise RuntimeError("too many conflicting saves")


def merge_index(a, b):
    out = {"v": 1, "items": dict((a or {}).get("items", {}))}
    for k, it in (b or {}).get("items", {}).items():
        if k not in out["items"] or it.get("updated", 0) > out["items"][k].get("updated", 0):
            out["items"][k] = it
    return out


# ---------- reading the material ----------
def save_image(img, path, width=1280):
    img = img.convert("RGB")
    if img.width > width:
        img = img.resize((width, round(img.height * width / img.width)))
    img.save(path, quality=75, optimize=True)


def read_documents(files, assets):
    """Slides and documents -> (slides list for the brief, handout text)."""
    from PIL import Image
    slides, handouts, n = [], [], 0
    for f in files:
        ext = f.suffix.lower()
        if ext == ".pdf":
            import pypdfium2 as pdfium
            step(f"Reading {f.name}…")
            doc = pdfium.PdfDocument(str(f))
            n += 1
            for i, page in enumerate(doc, 1):
                name = f"d{n}-slide-{i:03d}.jpg"
                save_image(page.render(scale=1.6).to_pil(), assets / name)
                slides.append({"file": f"assets/{name}", "label": f"{f.stem}, slide {i}", "text": page.get_textpage().get_text_range().strip()})
        elif ext == ".pptx":
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
            step(f"Reading {f.name}…")
            n += 1
            for i, slide in enumerate(Presentation(str(f)).slides, 1):
                texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                    texts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text)
                for j, pic in enumerate([sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE], 1):
                    name = f"d{n}-slide-{i:03d}-{j}.jpg"
                    try:
                        from io import BytesIO
                        save_image(Image.open(BytesIO(pic.image.blob)), assets / name)
                        slides.append({"file": f"assets/{name}", "label": f"{f.stem}, slide {i}, picture {j}", "text": ""})
                    except Exception:
                        pass
                slides.append({"file": None, "label": f"{f.stem}, slide {i}", "text": "\n".join(texts)})
        elif ext == ".docx":
            xml = zipfile.ZipFile(f).read("word/document.xml").decode("utf8", "ignore")
            handouts.append(f"### {f.name}\n" + re.sub(r"<[^>]+>", "", xml.replace("</w:p>", "\n")).strip())
        elif ext in (".txt", ".md", ".srt", ".vtt"):
            handouts.append(f"### {f.name}\n" + f.read_text(errors="ignore"))
        elif ext in IMAGE_EXT:
            n += 1
            name = f"img{n}.jpg"
            save_image(Image.open(f), assets / name)
            slides.append({"file": f"assets/{name}", "label": f.stem, "text": ""})
    return slides, "\n\n".join(handouts)


def mmss(t):
    t = int(t)
    return f"{t // 60:02d}:{t % 60:02d}"


def transcribe(media, work):
    import imageio_ffmpeg
    ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
    bindir = work / "bin"  # mlx-whisper calls `ffmpeg` by name
    bindir.mkdir(exist_ok=True)
    link = bindir / "ffmpeg"
    if not link.exists():
        link.symlink_to(ffmpeg)
    os.environ["PATH"] = f"{bindir}{os.pathsep}{os.environ.get('PATH', '')}"
    audio = work / "audio.wav"
    subprocess.run([ffmpeg, "-loglevel", "error", "-y", "-i", str(media), "-vn", "-ac", "1", "-ar", "16000", str(audio)], check=True)
    secs = audio.stat().st_size / 32000
    step(f"Transcribing {secs / 60:.0f} min of audio (roughly {max(1, secs / 600):.0f}–{max(2, secs / 300):.0f} min; the first time also downloads the speech model, ~1.6 GB)…")
    import mlx_whisper
    res = mlx_whisper.transcribe(str(audio), path_or_hf_repo="mlx-community/whisper-large-v3-turbo", condition_on_previous_text=False)
    audio.unlink(missing_ok=True)
    out, para, start = [], [], None
    for s in res["segments"]:
        start = s["start"] if start is None else start
        para.append(s["text"].strip())
        if s["end"] - start >= 30:
            out.append(f"[{mmss(start)}] " + " ".join(para))
            para, start = [], None
    if para:
        out.append(f"[{mmss(start)}] " + " ".join(para))
    return "\n\n".join(out)


def brief(job, slides, handouts, transcript):
    parts = ["## Lecture details", "\n".join(f"- {k}: {v}" for k, v in (
        ("title", job.get("title")), ("date", job.get("date")), ("course / branch", job.get("branch")),
        ("module / topic", job.get("topic"))) if v)]
    if slides:
        lines = []
        for im in slides:
            head = f"- {im['label']}" + (f" — embed as `{im['file']}`" if im.get("file") else " (text only)")
            text = (im.get("text") or "").strip().replace("\n", " / ")
            lines.append(head + (f"\n  Text on slide: {text[:600]}" if text else ""))
        parts += ["## Slides and images available", "Use the exact paths given when embedding images.", *lines]
    if handouts:
        parts += ["## Documents supplied with the lecture", handouts[:60000]]
    parts += ["## Transcript", transcript or "(No recording was supplied: write the notes from the slides and documents.)"]
    return "\n\n".join(parts)


def write_notes(work, engine, engine_path, model):
    prompt = (f"You are writing lecture notes. Read the style guide at {GUIDE} and follow it exactly, in English. "
              f"The lecture folder is the current directory. BRIEF.md has the details, the list of slide images (with their "
              f"exact paths under assets/), any documents, and the transcript. Look at the images listed in BRIEF.md with the "
              f"Read tool before writing (you may skip plain bullet-text slides whose text is given). There are no other "
              f"lectures to link to, so leave out cross-links except a short Connections note in prose. Then write the "
              f"complete notes to notes.md in this folder. Do not modify any other file.")
    env = {k: v for k, v in os.environ.items() if not k.startswith(("CLAUDECODE", "CLAUDE_CODE_")) and k != "ANTHROPIC_API_KEY"}
    if engine == "codex":
        cmd = [engine_path, "exec", "--skip-git-repo-check", "--sandbox", "workspace-write", "--color", "never", prompt]
    else:
        cmd = [engine_path, "-p", prompt, "--permission-mode", "acceptEdits", "--allowedTools", "Read", "Write", "Edit",
               "Glob", "Grep", "--add-dir", str(HERE), "--model", model or "opus"]
    step("Writing the notes with your " + ("ChatGPT" if engine == "codex" else "Claude") + " (this can take 5–15 minutes)…")
    p = subprocess.run(cmd, cwd=work, env=env, capture_output=True, text=True, timeout=3600)
    if not (work / "notes.md").exists():
        err = (p.stderr or p.stdout or "")[-600:]
        if re.search(r"log ?in|auth|credential|401", err, re.I):
            raise RuntimeError("Your AI tool isn't signed in: run `claude` (or `codex login`) once in Terminal.")
        raise RuntimeError("The AI didn't produce notes. " + err)
    return (work / "notes.md").read_text()


def main():
    work = Path(sys.argv[1])
    job = json.loads((work / "job.json").read_text())
    files = sorted((work / "files").iterdir())
    assets = work / "assets"
    assets.mkdir(exist_ok=True)
    space = Space(job["secret"], job["worker"])
    step("Checking your space…")
    space.get("space:index")  # fail early if the keys or connection are wrong

    recordings = [f for f in files if f.suffix.lower() in AUDIO_EXT]
    docs = [f for f in files if f not in recordings]
    slides, handouts = read_documents(docs, assets)
    transcript = ""
    for i, rec in enumerate(recordings):
        step(f"Preparing {rec.name}…")
        t = transcribe(rec, work)
        transcript += (f"\n\n## Recording {i + 1}: {rec.name}\n\n" if len(recordings) > 1 else "") + t
    if not transcript and not slides and not handouts:
        raise RuntimeError("Add a recording, slides or a document.")
    (work / "BRIEF.md").write_text(brief(job, slides, handouts, transcript))
    md = write_notes(work, job.get("engine", "claude"), job["engine_path"], job.get("model", "opus"))

    note_id = secrets.token_hex(12)
    used = sorted(set(re.findall(r"\]\((assets/[^)\s]+)\)", md)))
    for i, rel in enumerate(used, 1):
        f = work / rel
        if f.is_file():
            step(f"Saving images {i}/{len(used)}…")
            space.put(f"space:asset:{note_id}:{rel}", {"type": "image/jpeg", "b64": base64.b64encode(f.read_bytes()).decode()})
    step("Saving the notes to your space…")
    meta = {k: job.get(k, "") for k in ("title", "branch", "topic", "date")}
    space.put(f"space:note:{note_id}", {"v": 1, "md": md, "transcript": transcript,
                                       "slides": [{"label": s["label"], "text": s.get("text", "")} for s in slides], "meta": meta})
    t = int(time.time() * 1000)
    item = {"id": note_id, "kind": "note", **meta, "created": t, "updated": t}
    space.put("space:index", {"v": 1, "items": {note_id: item}}, merge=lambda theirs, mine: merge_index(theirs, mine))
    say(done=True, id=note_id, words=len(md.split()))


if __name__ == "__main__":
    try:
        main()
    except Exception as e:  # noqa: BLE001 - report everything to the page
        say(error=str(e))
        sys.exit(1)
